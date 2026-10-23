"""驗收測試（後端）。執行：python -m unittest discover -s tests -v
全部使用 tests/fixtures 的虛構資料。"""
import io, json, os, sqlite3, sys, tempfile, time, unittest
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import app as A
from btsys import db as dbm, options
from werkzeug.security import generate_password_hash
from docx import Document

FIX = os.path.join(os.path.dirname(__file__), 'fixtures')
H = {'X-BT-Requested': '1'}

class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.app = A.create_app({'DATABASE': os.environ.get('BT_TEST_DATABASE_URL') or os.path.join(cls.tmp, 'db.sqlite3'), 'UPLOAD_DIR': os.path.join(cls.tmp, 'up'), 'TESTING': True})
        with cls.app.app_context():
            d = dbm.get_db()
            if os.environ.get('BT_TEST_DATABASE_URL'):
                for t in ('audit', 'files', 'drafts', 'curriculum_imports', 'users'): d.execute(f'DELETE FROM {t}')
            for u in ('t1', 't2'):
                d.execute('INSERT INTO users(username,display_name,password_hash,is_admin,created_at) VALUES(?,?,?,0,?)',
                          (u, u, generate_password_hash('password12345'), time.time()))
            d.commit()

    def client(self, user='t1'):
        c = self.app.test_client()
        r = c.post('/api/login', json={'username': user, 'password': 'password12345'}, headers=H)
        self.assertEqual(r.status_code, 200)
        return c

    def new_plan(self, c, **basic):
        b = dict(dept='sh', grade='s12', shDept='餐飲服務科', domain='sn:社會技巧', subjectLabel='社會技巧', unit='上課好好坐', minutes=50)
        b.update(basic)
        st = [dict(key='s1', code='甲', name='甲'), dict(key='s2', code='乙', name='乙'), dict(key='s3', code='丙', name='丙')]
        return c.post('/api/drafts', json=dict(basic=b, students=st), headers=H).json['id']

    def upload(self, c, did, key, fn):
        return c.post(f'/api/drafts/{did}/files', data={'studentKey': key, 'file': (open(os.path.join(FIX, fn), 'rb'), fn)},
                      headers=H, content_type='multipart/form-data').json

class T1Options(Base):
    def test_four_departments_and_linked_domains(self):
        deps = [d['label'] for d in options.OPTIONS['departments']]
        self.assertEqual(deps, ['幼兒園', '國小部', '國中部', '技高部'])
        kg = [x['label'] for x in options.domains_for('kg')]
        self.assertIn('身體動作與健康', kg); self.assertNotIn('社會技巧', kg)
        e1 = [x['subject'] for x in options.domains_for('elem', 'e1')]
        self.assertIn('生活課程', e1); self.assertNotIn('社會', e1); self.assertIn('生活管理', e1)
        e5 = [x['subject'] for x in options.domains_for('elem', 'e5')]
        self.assertIn('社會', e5); self.assertNotIn('生活課程', e5)
        j7 = [x['subject'] for x in options.domains_for('jh', 'j7')]
        j8 = [x['subject'] for x in options.domains_for('jh', 'j8')]
        j9 = [x['subject'] for x in options.domains_for('jh', 'j9')]
        self.assertIn('科技', j7); self.assertIn('溝通訓練', j7); self.assertNotIn('溝通訓練', j8); self.assertIn('職業教育', j9)
        food = [x['label'] for x in options.domains_for('sh', 's10', '餐飲服務科')]
        for n in ('食材處理實作', '家事處理實作', '餐飲服務禮儀實作', '服務導論', '社會技巧'):
            self.assertIn(n, food)
        self.assertNotIn('居家點心實作', food); self.assertNotIn('飲食調理實作', food)
        home = [x['label'] for x in options.domains_for('sh', 's10', '居家生活服務科')]
        self.assertIn('居家點心實作', home); self.assertIn('飲食調理實作', home); self.assertNotIn('食材處理實作', home)
        self.assertEqual([d['key'] for d in options.find_dept('sh')['clusters'][0]['departments']], ['居家生活服務科', '餐飲服務科'])
        for d in ('kg', 'elem', 'jh', 'sh'):
            self.assertIn('其他／自訂', [x['label'] for x in options.domains_for(d)])

    def test_revalidation_after_switch(self):
        c = self.client()
        sel = [dict(itemId='SN-特社3-Ⅴ-4', code='特社3-Ⅴ-4')]
        r = c.post('/api/curriculum/validate', json=dict(context=dict(dept='kg', grade='kg-5', domain='kg:認知'), selection=sel), headers=H).json
        self.assertEqual(r['results'][0]['level'], 'error')
        r = c.post('/api/curriculum/validate', json=dict(context=dict(dept='elem', grade='e2', domain='sn:生活管理'), selection=sel), headers=H).json
        self.assertEqual(r['results'][0]['level'], 'warn')
        sv = [dict(itemId='SV-服務-技-車技Ⅰ-1', code='服務-技-車技Ⅰ-1')]
        r = c.post('/api/curriculum/validate', json=dict(context=dict(dept='sh', grade='s10', domain='sv:食材處理實作', shDept='餐飲服務科'), selection=sv), headers=H).json
        self.assertEqual(r['results'][0]['level'], 'error')
        r = c.post('/api/curriculum/validate', json=dict(context=dict(dept='jh', grade='j7', domain='sn:社會技巧'), selection=sv), headers=H).json
        self.assertEqual(r['results'][0]['level'], 'error')

class T2MultiStudent(Base):
    def test_separate_extraction_and_manual_edit(self):
        c = self.client(); did = self.new_plan(c)
        for k, fn in (('s1', '虛構學生甲_觀察紀錄.txt'), ('s2', '虛構學生乙_評量紀錄.txt'), ('s3', '虛構學生丙_IEP摘要.txt')):
            self.assertEqual(self.upload(c, did, k, fn)['files'][0]['status'], '已擷取')
        an = {k: c.post(f'/api/drafts/{did}/analyze', json={'studentKey': k}, headers=H).json['analysis'] for k in ('s1', 's2', 's3')}
        self.assertTrue(all(a['evidence']['filename'].startswith('虛構學生甲') for a in an['s1']['abilities']))
        self.assertTrue(all(a['evidence']['filename'].startswith('虛構學生乙') for a in an['s2']['abilities']))
        self.assertIn('第', an['s2']['abilities'][0]['evidence']['label'])
        plan = c.get(f'/api/drafts/{did}').json['plan']
        plan['analyses']['s2']['abilities'][0]['teacherText'] = '教師修正後的敘述'
        plan['analyses']['s2']['abilities'][0]['confirmed'] = True
        plan['students'][1]['abilitySummary'] = '教師撰寫摘要'
        c.put(f'/api/drafts/{did}', json=plan, headers=H)
        p2 = c.get(f'/api/drafts/{did}').json['plan']
        self.assertEqual(p2['analyses']['s2']['abilities'][0]['teacherText'], '教師修正後的敘述')
        self.assertEqual(p2['students'][1]['abilitySummary'], '教師撰寫摘要')

class T3Curriculum(Base):
    def test_separate_fields_and_traceable(self):
        c = self.client()
        r = c.post('/api/curriculum/recommend', json=dict(context=dict(dept='sh', grade='s12', domain='sn:社會技巧'), unit='上課好好坐'), headers=H).json
        self.assertTrue(r['core'] and r['performance'] and r['content'])
        self.assertTrue(all(x['type'] == '核心素養' for x in r['core']))
        self.assertTrue(all(x['type'] == '學習表現' for x in r['performance']))
        self.assertTrue(all(x['type'] == '學習內容' for x in r['content']))
        for x in r['performance'][:3]:
            self.assertTrue(x['source']['title'] and x['source']['version'] and x['source']['pdfPage'])
        # 待確認：無資料領域不給假資料
        r = c.post('/api/curriculum/recommend', json=dict(context=dict(dept='kg', grade='kg-5', domain='kg:認知'), unit='顏色分類'), headers=H).json
        self.assertEqual(r['status'], '待確認'); self.assertEqual(r['performance'], [])
        r = c.post('/api/curriculum/recommend', json=dict(context=dict(dept='elem', grade='e3', domain='gen:自然科學'), unit='植物'), headers=H).json
        self.assertEqual(r['status'], '待確認')
        # 官方原文被改寫 → 錯誤
        v = c.post('/api/curriculum/validate', json=dict(context=dict(dept='sh', grade='s12', domain='sn:社會技巧'),
                   selection=[dict(itemId='SN-特社3-Ⅴ-4', code='特社3-Ⅴ-4', officialText='改寫過的文字')]), headers=H).json
        self.assertEqual(v['results'][0]['level'], 'error')

    def test_school_course(self):
        c = self.client()
        r = c.post('/api/curriculum/recommend', json=dict(context=dict(dept='sh', grade='s10', shDept='餐飲服務科', domain='sch:餐飲服務禮儀實作'), unit='接待禮儀 迎賓及引導入座方式'), headers=H).json
        self.assertEqual(r['status'], '已載入'); self.assertTrue(all(x['framework'] == 'service-group' for x in r['performance']))
        r = c.post('/api/curriculum/recommend', json=dict(context=dict(dept='sh', grade='s10', shDept='餐飲服務科', domain='sch:地板滾球'), unit='規則'), headers=H).json
        self.assertEqual(r['status'], '待確認'); self.assertIn('健康與體育', r['reason'])
        dom = options.find_domain('sh', 'sch:餐飲服務禮儀實作')
        self.assertTrue(any(u['title'] == '接待禮儀' for u in dom['outline']['units']))

    def test_import(self):
        c = self.client()
        csv = 'code,type,subject,text,version,sourceTitle\nX-1,學習表現,數學,測試原文,版本1,測試來源\n,學習表現,數學,缺代碼,v,s\n'
        r = c.post('/api/curriculum/import', data={'title': '測試', 'version': 'v1', 'file': (io.BytesIO(csv.encode()), 'a.csv')}, headers=H, content_type='multipart/form-data').json
        self.assertEqual(r['imported'], 1); self.assertEqual(len(r['errors']), 1)

class T4Goals(Base):
    def test_goals_three_domains_with_assessment(self):
        c = self.client(); did = self.new_plan(c)
        for k, fn in (('s1', '虛構學生甲_觀察紀錄.txt'), ('s2', '虛構學生乙_評量紀錄.txt'), ('s3', '虛構學生丙_IEP摘要.txt')):
            self.upload(c, did, k, fn); c.post(f'/api/drafts/{did}/analyze', json={'studentKey': k}, headers=H)
        r = c.post(f'/api/drafts/{did}/goals/suggest', headers=H)
        self.assertEqual(r.status_code, 400)  # 未選課綱不產生目標
        plan = c.get(f'/api/drafts/{did}').json['plan']
        rec = c.post('/api/curriculum/recommend', json=dict(context=dict(dept='sh', grade='s12', domain='sn:社會技巧'), unit='上課好好坐'), headers=H).json
        pick = lambda x: dict(itemId=x['id'], code=x['code'], officialText=x['text'], teacherText='')
        plan['curriculum'] = dict(core=[pick(rec['core'][0])], performance=[pick(rec['performance'][0])], content=[pick(rec['content'][0])])
        c.put(f'/api/drafts/{did}', json=plan, headers=H)
        g = c.post(f'/api/drafts/{did}/groups/suggest', json={}, headers=H).json
        self.assertIn('障礙類別', g['note'])  # 說明不依障礙類別
        plan['groups'] = g['groups']; c.put(f'/api/drafts/{did}', json=plan, headers=H)
        goals = c.post(f'/api/drafts/{did}/goals/suggest', headers=H).json['goals']
        for k, lst in goals.items():
            self.assertEqual({x['domain'] for x in lst}, {'認知', '情意', '技能'})
            for x in lst:
                self.assertTrue(x['assessment'] and x['criterion'] and x['situation'] and x['support'] and x['behavior'])
                self.assertTrue(x['criterionSuggested'])
                self.assertTrue(x['performanceCodes'] and x['contentCodes'])
        # 乙無口語 → 回應方式含指認／溝通輔具
        grp_of_s2 = next(gg['key'] for gg in g['groups'] if 's2' in gg['studentKeys'])
        self.assertIn('指認', goals[grp_of_s2][0]['responseModes'])
        plan['goals'] = goals; c.put(f'/api/drafts/{did}', json=plan, headers=H)
        acts = c.post(f'/api/drafts/{did}/activities/suggest', headers=H).json
        self.assertEqual(acts['total'], 50)
        self.assertEqual([a['phase'] for a in acts['activities']], ['一、複習及引起動機', '二、發展活動', '三、綜合活動'])

class T5NoFabrication(Base):
    def test_missing_not_filled(self):
        c = self.client(); did = self.new_plan(c)
        self.upload(c, did, 's2', '虛構學生乙_評量紀錄.txt')
        a = c.post(f'/api/drafts/{did}/analyze', json={'studentKey': 's2'}, headers=H).json['analysis']
        self.assertEqual(a['basic']['gender']['value'], ''); self.assertEqual(a['basic']['gender']['status'], '資料不足')
        self.assertEqual(a['basic']['disability']['value'], '')
        self.assertIn('情緒行為', [x['domain'] for x in a['insufficient']])
        raw = open(os.path.join(FIX, '虛構學生乙_評量紀錄.txt'), encoding='utf-8').read()
        for ab in a['abilities']:
            self.assertIn(ab['quote'], raw)  # 每句能力描述都是原文
            self.assertEqual(ab['status'], '原文事實'); self.assertEqual(ab['relevance']['status'], '系統推估')
        # 無資料的學生不能分析
        r = c.post(f'/api/drafts/{did}/analyze', json={'studentKey': 's3'}, headers=H)
        self.assertEqual(r.status_code, 400)
        # 不支援格式明確告知
        r = c.post(f'/api/drafts/{did}/files', data={'studentKey': 's3', 'file': (io.BytesIO(b'xx'), 'a.doc')}, headers=H, content_type='multipart/form-data').json
        self.assertEqual(r['files'][0]['status'], '不支援'); self.assertIn('.docx', r['files'][0]['error'])
        # 無文字層的 PDF（模擬掃描檔）：沒有 OCR 時不產生任何內容
        from reportlab.pdfgen import canvas
        bio = io.BytesIO(); cv = canvas.Canvas(bio); cv.rect(50, 50, 200, 200); cv.save()
        r = c.post(f'/api/drafts/{did}/files', data={'studentKey': 's3', 'file': (io.BytesIO(bio.getvalue()), 'scan.pdf')}, headers=H, content_type='multipart/form-data').json
        f = r['files'][0]
        self.assertEqual(f['segments'], 0); self.assertTrue(any('掃描' in w or '未擷取' in w for w in f['warnings']))
        r = c.post(f'/api/drafts/{did}/analyze', json={'studentKey': 's3'}, headers=H)
        self.assertEqual(r.status_code, 400)

    def test_contradiction(self):
        c = self.client(); did = self.new_plan(c)
        c.post(f'/api/drafts/{did}/paste', json={'studentKey': 's1', 'text': '性別：男\n能理解指令。'}, headers=H)
        c.post(f'/api/drafts/{did}/paste', json={'studentKey': 's1', 'text': '性別：女'}, headers=H)
        a = c.post(f'/api/drafts/{did}/analyze', json={'studentKey': 's1'}, headers=H).json['analysis']
        self.assertEqual(a['basic']['gender']['status'], '資料矛盾'); self.assertEqual(a['basic']['gender']['value'], '')

class T6Export(Base):
    def test_docx_structure(self):
        c = self.client(); did = self.new_plan(c, dateText='115年10月30日', teacher='王老師', totalPeriods=2, periodNo=1, stageText='高三 302班')
        r = c.get(f'/api/drafts/{did}/export.docx'); self.assertEqual(r.status_code, 200)
        d = Document(io.BytesIO(r.data))
        self.assertEqual(len(d.tables), 3)
        t1, t2, t3 = d.tables
        self.assertEqual([c.text for c in t1.rows[0].cells], ['臺北市立臺北特殊教育學校', '□幼兒園□國小部\n□國中部█技高部', '教學設計表'])
        self.assertEqual(len(t2.columns), 9)  # 原檔九欄網格
        texts = [c.text for row in t2.rows for c in row.cells]
        for lbl in ('教學日期', '教學者', '領域/科目名稱', '單元名稱', '教學總節數', '本節為第__1___節', '教學階段\n(年級/班級/組別)',
                    '學\n生\n先備能力\n分\n析', '組別', '姓名', '性別', '障礙\n類別', '與本單元教學相關之能力', '課程\n架構', '核心\n素養', '學習\n重點', '議題\n融入', '學習\n目標'):
            self.assertIn(lbl, texts)
        sec = d.sections[0]
        self.assertAlmostEqual(sec.left_margin.pt, 34, delta=0.5); self.assertAlmostEqual(sec.top_margin.pt, 56.75, delta=0.5)
        self.assertIn('待確認', texts)  # 未選課綱時呈現待確認
        self.assertEqual([c.text for c in t3.rows[0].cells], ['教學活動', '教學\n資源', '教學\n時間', '評量\n方式\n(需列出A、B、C組評量方式)'])
        self.assertIn('w:tblHeader', t3.rows[0]._tr.xml)  # 跨頁重複標題
        self.assertIn('w:tblLayout w:type="fixed"', t2._tbl.xml)  # 固定欄寬避免溢出

    def test_pdf(self):
        import shutil
        if not shutil.which('soffice'): self.skipTest('未安裝 LibreOffice')
        c = self.client(); did = self.new_plan(c)
        r = c.get(f'/api/drafts/{did}/export.pdf'); self.assertEqual(r.status_code, 200); self.assertTrue(r.data.startswith(b'%PDF'))

class T8Retention(Base):
    def test_next_day_purge(self):
        c = self.client(); did = self.new_plan(c)
        self.upload(c, did, 's1', '虛構學生甲_觀察紀錄.txt'); self.upload(c, did, 's2', '虛構學生乙_評量紀錄.txt')
        with self.app.app_context():
            db = dbm.get_db()
            fid = db.execute('SELECT id FROM files WHERE draft_id=? ORDER BY id', (did,)).fetchone()['id']
            db.execute('UPDATE files SET created_at=? WHERE id=?', (time.time() - 86400 * 2, fid)); db.commit()
        os.environ['CRON_SECRET'] = 'test-cron'
        anon = self.app.test_client()
        self.assertEqual(anon.get('/api/cron/purge').status_code, 403)
        r = anon.get('/api/cron/purge', headers={'Authorization': 'Bearer test-cron'}).json
        self.assertGreaterEqual(r['deleted'], 1)
        files = c.get(f'/api/drafts/{did}').json['files']
        self.assertEqual(len(files), 1); self.assertEqual(files[0]['studentKey'], 's2')
        # 前一天建立的教案（含學生資料）整份刪除
        old = self.new_plan(c)
        self.upload(c, old, 's1', '虛構學生甲_觀察紀錄.txt')
        with self.app.app_context():
            db = dbm.get_db(); db.execute('UPDATE drafts SET created_at=? WHERE id=?', (time.time() - 86400 * 2, old)); db.commit()
        anon.get('/api/cron/purge', headers={'Authorization': 'Bearer test-cron'})
        self.assertEqual(c.get(f'/api/drafts/{old}').status_code, 404)
        self.assertEqual(c.get(f'/api/drafts/{did}').status_code, 200)

class T7Security(Base):
    def test_password_and_admin(self):
        c = self.client()
        r = c.post('/api/me/password', json={'old': 'wrong', 'new': 'abcdefghijk'}, headers=H); self.assertEqual(r.status_code, 400)
        r = c.post('/api/me/password', json={'old': 'password12345', 'new': 'password12345'}, headers=H); self.assertEqual(r.status_code, 200)
        self.assertEqual(c.get('/api/admin/users').status_code, 403)  # 非管理者

    def test_auth_isolation_csrf_encryption(self):
        anon = self.app.test_client()
        self.assertEqual(anon.get('/api/drafts').status_code, 401)
        c = self.client(); did = self.new_plan(c)
        self.upload(c, did, 's1', '虛構學生甲_觀察紀錄.txt')
        other = self.client('t2')
        self.assertEqual(other.get(f'/api/drafts/{did}').status_code, 404)
        self.assertEqual(other.get(f'/api/drafts/{did}/export.docx').status_code, 404)
        self.assertEqual(c.post('/api/drafts', json={}).status_code, 403)  # 缺少 CSRF 標頭
        with self.app.app_context():
            con = dbm.get_db()
            blob = b''.join(r['content_enc'] for r in con.execute('SELECT content_enc FROM drafts').fetchall())
            blobs = [r['blob_enc'] for r in con.execute('SELECT blob_enc FROM files').fetchall()]
        self.assertNotIn('上課好好坐'.encode(), blob)
        self.assertTrue(blobs)
        for b in blobs:
            self.assertNotIn('挫折容忍度'.encode(), b)
        # 刪除檔案
        fid = c.get(f'/api/drafts/{did}').json['files'][0]['id']
        self.assertEqual(c.delete(f'/api/drafts/{did}/files/{fid}', headers=H).status_code, 200)
        self.assertEqual(c.get(f'/api/drafts/{did}').json['files'], [])
        self.assertEqual(c.delete(f'/api/drafts/{did}', headers=H).status_code, 200)

if __name__ == '__main__':
    unittest.main()
