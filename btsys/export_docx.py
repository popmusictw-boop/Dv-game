"""匯出 Word：依北特「1221公開授課教案格式.docx」的表格結構重建。

原檔結構（由原檔文字結構讀取）：
  表1（3欄）：臺北市立臺北特殊教育學校｜□幼兒園□國小部□國中部□技高部｜教學設計表
  表2（6欄）：教學日期／教學者、領域/科目名稱／單元名稱、教學總節數／本節為第＿節／教學階段、
             學生先備能力分析（左欄垂直合併；組別、姓名、性別、障礙類別、與本單元教學相關之能力）、
             課程架構（整列合併）、核心素養、學習重點（學習表現、學習內容）、議題融入、學習目標
  表3（4欄）：教學活動｜教學資源｜教學時間｜評量方式(需列出A、B、C組評量方式)
欄寬、列高、字型與邊界依 Google 文件 API 讀取原檔之版面數值重建（見 build_docx 說明）。
防破版：表格固定欄寬、文字於儲存格內自動換行；標題列跨頁重複；短列不跨頁分割。
"""
import io
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_EA = '標楷體'
FONT_LATIN = 'Times New Roman'
DEPTS = [('kg', '幼兒園'), ('elem', '國小部'), ('jh', '國中部'), ('sh', '技高部')]

def _font(run, size=11, bold=False):
    run.font.name = FONT_LATIN
    run.font.size = Pt(size); run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rf.set(qn(a), FONT_EA if a == 'w:eastAsia' else FONT_LATIN)

def _set_cell(cell, text, size=11, bold=False, align='left', valign='center', shade=None):
    cell.text = ''
    lines = (text or '').split('\n')
    p = cell.paragraphs[0]
    for i, line in enumerate(lines):
        if i: p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
        _font(p.add_run(line), size, bold)
    cell.vertical_alignment = {'center': WD_CELL_VERTICAL_ALIGNMENT.CENTER, 'top': WD_CELL_VERTICAL_ALIGNMENT.TOP, 'bottom': WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}[valign]
    if shade:
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), shade); tcPr.append(shd)

def _fixed_table(doc, rows, widths_cm, borders=True):
    t = doc.add_table(rows=rows, cols=len(widths_cm))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    if borders:
        t.style = 'Table Grid'
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths_cm[i] * 567)))
    return t

def _no_split(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement('w:cantSplit'); trPr.append(e)

def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement('w:tblHeader'); trPr.append(e)

def _page_number_footer(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(); _font(run, 10)
    for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tag:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), tag); run._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; run._r.append(it)

def plan_fields(plan):
    """把教案資料整理為各儲存格文字（預覽與 Word 共用同一份）。"""
    b = plan.get('basic', {})
    cur = plan.get('curriculum', {})
    def items(lst):
        return '\n'.join(f'{x["code"]} {x.get("teacherText") or x["officialText"]}' for x in lst) or '待確認'
    groups = plan.get('groups', [])
    goals = plan.get('goals', {})
    gl = []
    for g in groups:
        gg = goals.get(g['key']) or []
        if not g.get('studentKeys') and not gg: continue
        gl.append(f'{g["name"]}目標')
        for i, x in enumerate(gg, 1):
            gl.append(f'{i}.（{x["domain"]}）{x.get("text") or ""}')
    return dict(
        dept=b.get('dept'), date=b.get('dateText') or b.get('date') or '', teacher=b.get('teacher', ''),
        subject=b.get('subjectLabel', ''), unit=b.get('unit', ''), totalPeriods=str(b.get('totalPeriods') or ''),
        periodNo=str(b.get('periodNo') or ''), stageText=b.get('stageText', ''),
        core=items(cur.get('core', [])),
        focus='學習表現：\n' + items(cur.get('performance', [])) + '\n學習內容：\n' + items(cur.get('content', [])),
        issues=plan.get('issuesText') or '', goals='\n'.join(gl) or '待確認')

PT = 1 / 28.3465  # pt → cm

def _row_height(row, pt):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement('w:trHeight')
    e.set(qn('w:val'), str(int(pt * 20))); e.set(qn('w:hRule'), 'atLeast'); trPr.append(e)

def build_docx(plan):
    """依原檔（Google 文件 API 讀取之版面）重建：A4、上下邊界 56.75/56.7pt、左 34pt、右 28.35pt；
    字型標楷體（DFKai-SB）12pt；表1 三欄 180.45/157.2/106.3pt；表2 九欄網格 20.9/22.5/0.35/48/21.3/63.7/93.15/148.85/127.55pt；
    表3 四欄 383.65/48.7/50.5/62.9pt；各列最小列高同原檔。"""
    f = plan_fields(plan)
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Pt(56.75); sec.bottom_margin = Pt(56.7); sec.left_margin = Pt(34); sec.right_margin = Pt(28.35)
    sec.header_distance = Pt(42.55); sec.footer_distance = Pt(49.6)
    _page_number_footer(sec)
    st = doc.styles['Normal']; st.font.size = Pt(12); st.font.name = FONT_LATIN
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_EA)
    st.paragraph_format.space_after = Pt(0)
    S = 12  # 原檔內文字級

    # 表1：抬頭（臺北市立臺北特殊教育學校｜部別｜教學設計表）
    t1 = _fixed_table(doc, 1, [180.45 * PT, 157.2 * PT, 106.3 * PT])
    boxes = ''.join(('█' if k == f['dept'] else '□') + lbl + ('\n' if k == 'elem' else '') for k, lbl in DEPTS)
    _set_cell(t1.cell(0, 0), '臺北市立臺北特殊教育學校', 20, True, 'center')
    _set_cell(t1.cell(0, 1), boxes, 16, True, 'center')
    _set_cell(t1.cell(0, 2), '教學設計表', 20, True, 'center')
    mar = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), '20'); e.set(qn('w:type'), 'dxa'); mar.append(e)
    t1._tbl.tblPr.append(mar)
    doc.add_paragraph()

    # 表2：九欄網格
    order = {k: i for i, g in enumerate(plan.get('groups', [])) for k in g.get('studentKeys', [])}
    students = sorted([s for s in plan.get('students', []) if not s.get('excludeFromPlan')], key=lambda s: order.get(s.get('key'), 99)) or [{}]
    W = [20.9, 22.5, 0.35, 48, 21.3, 63.7, 93.15, 148.85, 127.55]
    t2 = _fixed_table(doc, 4 + len(students) + 5, [w * PT for w in W])
    c = lambda r, a, b=None: t2.cell(r, a) if b is None else t2.cell(r, a).merge(t2.cell(r, b))
    lab = lambda cell, text, va='center': _set_cell(cell, text, S, False, 'center', va)
    val = lambda cell, text, al='center', va='center', size=S: _set_cell(cell, text, size, False, al, va)
    for r, (l1, v1, l2, v2) in enumerate((('教學日期', f['date'], '教學者', f['teacher']), ('領域/科目名稱', f['subject'], '單元名稱', f['unit']))):
        _row_height(t2.rows[r], 28.35); _no_split(t2.rows[r])
        lab(c(r, 0, 3), l1); val(c(r, 4, 6), v1); lab(c(r, 7), l2); val(c(r, 8), v2)
    r = 2; _no_split(t2.rows[r])
    lab(c(r, 0, 3), '教學總節數'); val(c(r, 4), f['totalPeriods'])
    val(c(r, 5, 6), f'本節為第__{f["periodNo"] or "_"}___節')
    lab(c(r, 7), '教學階段\n(年級/班級/組別)'); val(c(r, 8), f['stageText'])
    r = 3; _row_height(t2.rows[r], 31.7); _no_split(t2.rows[r])
    lab(c(r, 1, 2), '組別'); lab(c(r, 3), '姓名'); lab(c(r, 4), '性別'); lab(c(r, 5), '障礙\n類別'); lab(c(r, 6, 8), '與本單元教學相關之能力')
    gname = {g['key']: g['name'].replace('組', '') for g in plan.get('groups', [])}
    gof = {k: g['key'] for g in plan.get('groups', []) for k in g.get('studentKeys', [])}
    for i, s in enumerate(students):
        r = 4 + i; _row_height(t2.rows[r], 28.35)
        val(c(r, 1, 2), gname.get(gof.get(s.get('key')), ''), 'center', 'top'); val(c(r, 3), s.get('name') or s.get('code') or '', 'center', 'top')
        val(c(r, 4), s.get('gender', ''), 'center', 'top'); val(c(r, 5), s.get('disability', ''), 'center', 'top')
        val(c(r, 6, 8), s.get('abilitySummary', '') or '待確認', 'left', 'top')
    lab(t2.cell(3, 0).merge(t2.cell(3 + len(students), 0)), '學\n生\n先備能力\n分\n析')
    r = 4 + len(students)
    for k, (lbl, v) in enumerate((('課程\n架構', plan.get('courseStructure', '')), ('核心\n素養', f['core']), ('學習\n重點', f['focus']),
                                   ('議題\n融入', f['issues'] or '無'), ('學習\n目標', f['goals']))):
        row = t2.rows[r + k]; _row_height(row, 52.5)
        if k == 0: _no_split(row)
        lab(c(r + k, 0, 1), lbl); val(c(r + k, 2, 8), v, 'left', 'top')

    doc.add_paragraph()
    # 表3：教學流程（單一內容列，各欄以階段序號對齊）
    acts = plan.get('activities', [])
    t3 = _fixed_table(doc, 2, [383.65 * PT, 48.7 * PT, 50.5 * PT, 62.9 * PT])
    hd = t3.rows[0]; _repeat_header(hd); _no_split(hd)
    lab(hd.cells[0], '教學活動'); lab(hd.cells[1], '教學\n資源'); _set_cell(hd.cells[2], '教學\n時間', 11, False, 'center')
    _set_cell(hd.cells[3], '評量\n方式\n(需列出A、B、C組評量方式)', 9, False, 'center')
    content, res, tm, ass = [], [], [], []
    for a in acts:
        content.append(a['phase']); content.append(a.get('content', ''))
        gs = a.get('groupSupport') or {}
        if any(gs.values()):
            content.append('【各組支持】' + '；'.join(f'{gname.get(k, k)}組：{v}' for k, v in gs.items() if v))
        mark = a['phase'].split('、')[0] + '、'
        res.append(mark + (a.get('resources') or '')); tm.append(f'{a.get("minutes", "")}')
        asg = a.get('assessment') or {}
        ass.append(mark + '\n'.join(f'{gname.get(k, k)}:{v}' for k, v in asg.items() if v))
    row = t3.rows[1]
    val(row.cells[0], '\n'.join(content), 'left', 'top')
    val(row.cells[1], '\n\n'.join(res), 'justify', 'top')
    val(row.cells[2], '\n\n'.join(tm), 'center', 'top', 11)
    val(row.cells[3], '\n\n'.join(ass), 'justify', 'top', 10)
    bio = io.BytesIO(); doc.save(bio); return bio.getvalue()
