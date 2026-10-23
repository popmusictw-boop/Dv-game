"""學生資料檔案文字擷取：PDF、Word（.docx）、圖片、純文字。

每個段落保留來源定位（頁碼、段落序號），作為後續「原始資料佐證」。
掃描 PDF 與圖片使用 Tesseract OCR；伺服器未安裝中文辨識資料時，明確回報無法辨識，不猜測內容。
"""
import io, re, shutil, subprocess

SUPPORTED = {
    'pdf': 'PDF（文字型與掃描型；掃描型需伺服器 OCR）',
    'docx': 'Word（.docx）',
    'png': '圖片', 'jpg': '圖片', 'jpeg': '圖片', 'bmp': '圖片', 'tif': '圖片', 'tiff': '圖片', 'webp': '圖片',
    'txt': '純文字（UTF-8）',
}
UNSUPPORTED_HINT = {
    'doc': '舊版 Word（.doc）無法可靠擷取，請在 Word 另存為 .docx 或 PDF 後再上傳。',
    'heic': 'iPhone HEIC 相片格式不支援，請在手機設定「相機→格式→最相容」或轉成 JPG 後上傳。',
    'pages': 'Pages 文件請匯出為 PDF 或 Word（.docx）。',
    'odt': 'ODT 請另存為 .docx 或 PDF。',
    'xls': '試算表請另存為 PDF，或把相關內容貼上為文字。', 'xlsx': '試算表請另存為 PDF，或把相關內容貼上為文字。',
}

class Unsupported(Exception):
    pass

def ext_of(name):
    return name.rsplit('.', 1)[-1].lower() if '.' in name else ''

def ocr_status(lang):
    """回報 OCR 可用性（Tesseract 主程式與所需語言資料）。"""
    if not shutil.which('tesseract'):
        return dict(available=False, message='伺服器未安裝 Tesseract OCR，掃描檔與圖片無法辨識文字。')
    try:
        langs = subprocess.run(['tesseract', '--list-langs'], capture_output=True, text=True, timeout=10).stdout.split()
    except Exception:
        langs = []
    need = [l for l in lang.split('+')]
    missing = [l for l in need if l not in langs]
    if 'chi_tra' in missing:
        return dict(available=False, langs=langs, message='伺服器 OCR 缺少繁體中文辨識資料（chi_tra），掃描檔與圖片中的中文無法辨識。請管理者安裝 tesseract-ocr-chi-tra。')
    return dict(available=True, langs=langs, lang='+'.join(l for l in need if l in langs), message='OCR 可用')

def _split_paragraphs(text):
    paras = [p.strip() for p in re.split(r'\n\s*\n|\r\n\s*\r\n', text) if p.strip()]
    if len(paras) <= 1:
        paras = [p.strip() for p in text.splitlines() if p.strip()]
    return paras

def _ocr_image(img, lang):
    import pytesseract
    data = pytesseract.image_to_data(img, lang=lang, config='--psm 6', output_type=pytesseract.Output.DICT)
    lines, confs = {}, {}
    for i, w in enumerate(data['text']):
        if not w.strip(): continue
        key = (data['block_num'][i], data['par_num'][i], data['line_num'][i])
        lines.setdefault(key, []).append(w)
        c = float(data['conf'][i]);
        if c >= 0: confs.setdefault(key, []).append(c)
    out = []
    for key in sorted(lines):
        t = ''.join(lines[key]) if 'chi' in lang else ' '.join(lines[key])
        c = confs.get(key) or [0]
        out.append((t, sum(c) / len(c)))
    return out

def extract(data: bytes, filename: str, ocr_lang='chi_tra+eng'):
    """回傳 dict(segments=[{loc, text, conf?}], method, warnings, pages)。"""
    ext = ext_of(filename)
    if ext in UNSUPPORTED_HINT:
        raise Unsupported(UNSUPPORTED_HINT[ext])
    if ext not in SUPPORTED:
        raise Unsupported(f'不支援的檔案格式（.{ext or "無副檔名"}）。支援：PDF、Word（.docx）、JPG／PNG 等圖片、純文字，或直接貼上文字。')
    segs, warnings, method, pages = [], [], '', None
    ocr = ocr_status(ocr_lang)
    if ext == 'pdf':
        import pdfplumber
        method = 'PDF 文字層'
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = len(pdf.pages)
            scanned_pages = []
            for pno, page in enumerate(pdf.pages, 1):
                txt = page.extract_text() or ''
                if len(re.sub(r'\s', '', txt)) < 15:
                    scanned_pages.append(pno); continue
                for k, p in enumerate(_split_paragraphs(txt), 1):
                    segs.append(dict(loc=dict(page=pno, para=k), text=p))
        if scanned_pages:
            if ocr['available']:
                import pypdfium2 as pdfium
                doc = pdfium.PdfDocument(data)
                method += '＋OCR'
                for pno in scanned_pages:
                    img = doc[pno - 1].render(scale=300 / 72).to_pil()
                    for k, (t, c) in enumerate(_ocr_image(img, ocr['lang']), 1):
                        segs.append(dict(loc=dict(page=pno, para=k, ocr=True), text=t, conf=round(c)))
            else:
                warnings.append(f'第 {"、".join(map(str, scanned_pages))} 頁為掃描影像或無文字層：{ocr["message"]}此頁內容未擷取，不會被當作學生資料。')
    elif ext == 'docx':
        import docx
        method = 'Word 段落與表格'
        d = docx.Document(io.BytesIO(data))
        k = 0
        for p in d.paragraphs:
            if p.text.strip():
                k += 1; segs.append(dict(loc=dict(para=k), text=p.text.strip()))
        for ti, tb in enumerate(d.tables, 1):
            for ri, row in enumerate(tb.rows, 1):
                cells, seen = [], set()
                for c in row.cells:
                    if id(c._tc) in seen: continue
                    seen.add(id(c._tc)); t = c.text.strip()
                    if t: cells.append(t)
                if cells:
                    segs.append(dict(loc=dict(table=ti, row=ri), text='｜'.join(cells)))
    elif ext == 'txt':
        method = '純文字'
        text = data.decode('utf-8-sig', errors='replace')
        for k, p in enumerate(_split_paragraphs(text), 1):
            segs.append(dict(loc=dict(para=k), text=p))
    else:
        method = '圖片 OCR'
        if not ocr['available']:
            raise Unsupported(ocr['message'] + '請改上傳文字型 PDF／Word，或直接貼上文字。')
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        for k, (t, c) in enumerate(_ocr_image(img, ocr['lang']), 1):
            segs.append(dict(loc=dict(page=1, para=k, ocr=True), text=t, conf=round(c)))
    low = [s for s in segs if s.get('conf') is not None and s['conf'] < 60]
    if low:
        warnings.append(f'有 {len(low)} 段 OCR 辨識信心偏低（<60），已標示，請對照原檔確認。')
    if not segs:
        warnings.append('未擷取到任何文字，請確認檔案內容或改以貼上文字方式提供。')
    return dict(segments=segs, method=method, warnings=warnings, pages=pages)

def from_pasted(text):
    return dict(segments=[dict(loc=dict(para=k), text=p) for k, p in enumerate(_split_paragraphs(text), 1)],
                method='教師貼上文字', warnings=[], pages=None)

def loc_label(filename, loc):
    parts = [filename]
    if loc.get('page'): parts.append(f'第{loc["page"]}頁')
    if loc.get('table'): parts.append(f'表格{loc["table"]}第{loc["row"]}列')
    if loc.get('para'): parts.append(f'第{loc["para"]}段')
    if loc.get('ocr'): parts.append('（OCR）')
    return ' '.join(parts)
